from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core import api_status
from app.core.config import settings
from app.core.exceptions import BusinessException
from app.core.security import validate_access_token
from app.integrations.iqms import IQMSClient
from app.main import app
from app.modules.inventory.service import WEBAPI_UNREADABLE_AKA_FIELDS, InventoryService
from app.modules.reporting.client_report import (
    AKA_UNREADABLE_READ_CODES,
    NOT_READABLE,
    UPDATE,
    build_quote_report,
)
from app.modules.reporting.client_workbook import _unique_sheet_name, build_client_workbook
from app.modules.reporting.schemas import (
    AkaResolution,
    ExceptionDetail,
    LineItemSummary,
    PricingResolution,
    PricingTier,
    QuoteSummaryData,
    SalesOrderFact,
    SalesOrderResolution,
)
from app.modules.sales_order.service import SalesOrderService, _build_default_iqms_client

client = TestClient(app)


@pytest.fixture(autouse=True)
def authenticated():
    app.dependency_overrides[validate_access_token] = lambda: {"sub": "test-user-id", "tenant": "test-tenant"}
    yield
    app.dependency_overrides.clear()


# ---------------------------------------------------------------------------
# Data-source register
# ---------------------------------------------------------------------------

def _openapi_operations():
    spec = app.openapi()
    for path, operations in spec["paths"].items():
        for method, operation in operations.items():
            if method in ("get", "post", "put", "patch", "delete"):
                yield method.upper(), path, operation


def test_every_served_operation_is_classified():
    served = {(method, path) for method, path, _ in _openapi_operations()}
    assert served == set(api_status.API_STATUS), "register and served operations differ"


def test_operation_summaries_carry_the_data_source_marker():
    for method, path, operation in _openapi_operations():
        assert operation["summary"].startswith(f"[{api_status.lookup(method, path).status}] ")
        assert "[UNCLASSIFIED]" not in operation["summary"]


def test_summary_marks_dummy_create_and_update_as_not_integrated():
    operations = {(m, p): op for m, p, op in _openapi_operations()}
    for path in ("/inventory/create-aka/dummy", "/inventory/update-aka/dummy", "/inventory/add-pricebreak/dummy",
                 "/inventory/update-pricebreak/dummy", "/inventory/get-pricebreaks/dummy",
                 "/sales-orders/update-sales-order-details/dummy"):
        operation = operations[("POST", path)]
        assert operation["summary"].startswith("[DUMMY] ")
        assert operation["x-data-source"]["integration"] == api_status.NOT_INTEGRATED


def test_webapi_backed_endpoints_are_live():
    for path in ("/inventory/get-aka", "/inventory/get-pricebreaks", "/sales-orders/get-sales-orders",
                 "/sales-orders/get-sales-order-details", "/sales-orders/get-sales-order-releases",
                 "/customers/resolve"):
        assert api_status.lookup("POST", path).status == api_status.LIVE
        assert api_status.lookup("POST", path).integration == api_status.WEBAPI_READ


def test_response_header_names_the_data_source():
    response = client.post("/inventory/get-pricebreaks/dummy", json={})
    assert response.headers["X-Data-Source"] == api_status.DUMMY
    assert response.headers["X-Data-Source-Integration"] == api_status.NOT_INTEGRATED


def test_api_status_endpoint_lists_the_register():
    response = client.get("/api-status")
    assert response.status_code == 200
    rows = {(row["method"], row["path"]): row for row in response.json()["data"]}
    assert rows[("POST", "/inventory/create-aka/dummy")]["status"] == "DUMMY"
    assert rows[("POST", "/inventory/get-aka")]["status"] == "LIVE"
    assert len(rows) == len(api_status.API_STATUS)


# ---------------------------------------------------------------------------
# Live WebAPI wiring
# ---------------------------------------------------------------------------

def test_sales_order_default_client_is_live_unless_fixture_mode(monkeypatch):
    monkeypatch.setattr(settings, "IQMS_DATA_SOURCE", "live")
    assert isinstance(_build_default_iqms_client(), IQMSClient)

    monkeypatch.setattr(settings, "IQMS_DATA_SOURCE", "fixture")
    assert type(_build_default_iqms_client()).__name__ == "FixtureIQMSClient"


def test_sales_order_persists_records_with_iqms_source(monkeypatch):
    monkeypatch.setattr(settings, "IQMS_DATA_SOURCE", "live")
    fake_iqms = MagicMock()
    fake_iqms.get_sales_orders.return_value = [{
        "Id": 1, "OrdDetailId": 2, "ArInvtId": 3, "ItemNumber": "EVCO-1",
        "CustomerNumber": "CUST-1", "TotalQTYOrdered": 10, "UnitPrice": 2.5,
    }]
    sync_repository = MagicMock()
    pricing_history = MagicMock()
    pricing_history.get_active_price_for_quantity.return_value = None
    service = SalesOrderService(
        iqms_client=fake_iqms,
        sync_repository=sync_repository,
        pricing_history_repository=pricing_history,
        resolution_repository=MagicMock(),
    )

    orders = service.get_sales_orders(item_number="EVCO-1")

    assert len(orders) == 1
    assert sync_repository.upsert_sales_order.call_args.kwargs["source"] == "iqms"


def _aka_service(records, customer_found=True, processing_id=None):
    iqms = MagicMock()
    iqms.get_aka_inventory_for_customer.return_value = records
    customers = MagicMock()
    customers.find_candidates.return_value = (
        [SimpleNamespace(ar_custo_id=5, customer_name="Acme Corp")] if customer_found else []
    )
    processing = MagicMock()
    processing.get_processing_id_for_line_item.return_value = processing_id
    resolution = MagicMock()
    service = InventoryService(
        iqms_client=iqms, customer_repository=customers,
        resolution_repository=resolution, processing_repository=processing,
    )
    return service, resolution, processing


def test_get_aka_search_maps_webapi_record_and_leaves_unreadable_fields_empty():
    service, resolution, _ = _aka_service([{
        "CustomerNumber": "CUST-4471", "AKAItemNumber": "38795", "ItemNumber": "EVCO-10023",
        "AKADescription": "Bracket", "ItemDescription": "Bracket assembly",
    }])

    result = service.get_aka_search("EVCO-10023", "CUST-4471", "MFG-8890")

    assert result.header.item_number == "EVCO-10023"
    assert result.header.description == "Bracket assembly"
    detail = result.aka_details[0]
    assert detail.aka_item_number == "38795"
    assert detail.customer_name == "Acme Corp"
    assert detail.manufacturing_bom_number == ""  # not returned by WebAPI, never echoed from the request
    assert detail.currency == ""
    assert resolution.record_resolution.call_args.kwargs["decision"] == "NO_CHANGE"


def test_get_aka_response_contains_only_fields_the_webapi_returns(monkeypatch):
    from app.modules.inventory import router as inventory_router
    from app.modules.inventory.schemas import AkaDetail, AkaHeader, AkaSearchData

    result = AkaSearchData(
        header=AkaHeader(item_number="9475226", rev="0-04", description="Upper Door"),
        aka_details=[AkaDetail(
            aka_item_number="0017483-524", aka_description="Upper Door Entec", rev="A",
            customer_number="11762", customer_name="SANMINA", manufacturing_bom_number="",
            ship_to_attn="Dock 4", currency="", minimum_selling_qty=0, selling_multiples_of=0, so_item_number="",
        )],
    )
    monkeypatch.setattr(inventory_router.inventory_service, "get_aka_search", lambda **kwargs: result)
    monkeypatch.setattr(settings, "IQMS_DATA_SOURCE", "live")

    response = client.post("/inventory/get-aka", json={
        "evco_part_number": "9475226", "customer_number": "11762", "manufacturing_bom_number": "x",
    })

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["header"] == {"evco_part_number": "9475226", "rev": "0-04", "part_description": "Upper Door"}
    assert data["aka_details"] == [{
        "customer_part_number": "0017483-524", "part_description": "Upper Door Entec", "rev": "A",
        "customer_number": "11762", "customer_name": "SANMINA", "ship_to_attn": "Dock 4",
    }]


def test_get_aka_search_without_a_match_returns_header_only_and_records_exception():
    line_item_id, processing_id = uuid4(), uuid4()
    service, resolution, processing = _aka_service([], processing_id=processing_id)

    result = service.get_aka_search("EVCO-10023", "CUST-4471", "MFG-8890", line_item_id=line_item_id)

    assert result.aka_details == []
    assert resolution.record_resolution.call_args.kwargs["decision"] == "NOT_FOUND"
    assert processing.record_exception.call_args.kwargs["exception_code"] == "AKA_NOT_FOUND"


def test_get_price_breaks_maps_webapi_records_using_aka_ids():
    iqms = MagicMock()
    iqms.get_aka_inventory_for_customer.return_value = [
        {"ItemNumber": "9475226", "ArInvtId": 227097, "ArCustoId": 26166, "AKAItemNumber": "0017483-524"},
    ]
    iqms.get_aka_item_price_breaks.return_value = [
        {"Id": 21143, "UnitPrice": 31.63, "Quantity": 240, "Comment": "65150-003"},
    ]
    customers = MagicMock()
    customers.find_candidates.return_value = [SimpleNamespace(ar_custo_id=26166, customer_name="Sanmina")]
    service = InventoryService(iqms_client=iqms, customer_repository=customers,
                               resolution_repository=MagicMock(), processing_repository=MagicMock())

    result = service.get_price_breaks("11762", "9475226")

    iqms.get_aka_item_price_breaks.assert_called_once_with(227097, 26166)
    assert [(b.unit_price, b.quantity, b.comment) for b in result] == [(31.63, 240, "65150-003")]


def test_get_price_breaks_without_aka_record_returns_empty_list_and_skips_webapi_call():
    iqms = MagicMock()
    iqms.get_aka_inventory_for_customer.return_value = []
    customers = MagicMock()
    customers.find_candidates.return_value = [SimpleNamespace(ar_custo_id=26166, customer_name="Sanmina")]
    service = InventoryService(iqms_client=iqms, customer_repository=customers,
                               resolution_repository=MagicMock(), processing_repository=MagicMock())

    assert service.get_price_breaks("11762", "9475226") == []
    iqms.get_aka_item_price_breaks.assert_not_called()


def test_get_aka_search_unknown_customer_raises_404():
    service, _, _ = _aka_service([], customer_found=False)
    with pytest.raises(BusinessException) as excinfo:
        service.get_aka_search("EVCO-10023", "NOPE", "MFG-8890")
    assert excinfo.value.code == "CUSTOMER_NOT_FOUND"
    assert excinfo.value.status_code == 404


# ---------------------------------------------------------------------------
# Client report rows
# ---------------------------------------------------------------------------

def _line(**overrides):
    values = dict(
        line_item_id=uuid4(), line_number=1, evco_part_id="EVCO-10023", customer_part_number="CUST-4471",
        aka=AkaResolution(decision="NO_CHANGE"), pricing=PricingResolution(decision=None, tiers=[]),
        sales_order=SalesOrderResolution(decision=None, orders=[]), exceptions=[], exception_details=[],
    )
    values.update(overrides)
    return LineItemSummary(**values)


def _summary(*lines, quote_exceptions=()):
    return QuoteSummaryData(
        processing_id=uuid4(), quote_number="SAMPLE-Q-001", customer_name="Acme Corp", ar_cust_id="CUST-4471",
        line_items=list(lines), quote_exceptions=list(quote_exceptions),
    )


def test_changed_aka_pricing_sales_order_and_exception_rows():
    line_id = uuid4()
    line = _line(
        line_item_id=line_id,
        aka=AkaResolution(
            decision="UPDATED", aka_id="38795",
            before={"akaItem#": "38795", "mfg#": "OLD", "currency": "", "minimumSellingQty": 0},
            after={"akaItem#": "38795", "mfg#": "6601/9480026", "currency": "USD", "minimumSellingQty": 100},
        ),
        pricing=PricingResolution(decision="UPDATED", tiers=[PricingTier(
            quantity=500, previous_price=2.5, current_price=2.75, is_active=True,
        )]),
        sales_order=SalesOrderResolution(decision="MISMATCH", orders=[SalesOrderFact(
            sales_order_id="SO-45871", decision="MISMATCH",
            note="Order unit price $1.20 differs from current quoted price $1.45.",
        )]),
        exception_details=[ExceptionDetail(exception_code="PRICE_NOT_FOUND", message="none on file")],
    )

    report = build_quote_report(_summary(line))
    rows = {row.field: row for row in report.rows}

    mfg = rows["AKA manufacturing number (BOM)"]
    assert (mfg.current_value, mfg.proposed_value, mfg.rule) == (NOT_READABLE, "6601/9480026", "VR-009")
    assert mfg.comments == "AKA_READ_MFG; AKA_WRITE_GAP"
    assert rows["AKA minimum sell qty"].current_value == NOT_READABLE
    assert "AKA customer part number" not in rows  # unchanged

    price = rows["Unit price, qty tier 500"]
    assert (price.current_value, price.proposed_value, price.rule) == ("2.50", "2.75", "VR-016")

    so = rows["Unit price on SO SO-45871"]
    assert (so.current_value, so.proposed_value, so.rule, so.category) == ("$1.20", "$1.45", "VR-021", "Sales Order")

    assert rows["AKA currency"].current_value == NOT_READABLE
    assert rows["Price breaks"].rule == "VR-015"
    assert all(row.status == UPDATE for row in report.rows)
    # mfg#, currency, minimum sell qty, price tier, sales order, price-break exception
    assert report.total_rows == 6 and report.update_rows == 6 and report.no_action_rows == 0


def test_not_found_aka_is_reported_once_even_with_matching_exception():
    line = _line(
        aka=AkaResolution(decision="NOT_FOUND"),
        exception_details=[ExceptionDetail(exception_code="AKA_NOT_FOUND")],
    )
    report = build_quote_report(_summary(line))
    assert len(report.rows) == 1
    assert report.rows[0].comments == "AKA_NOT_FOUND; AKA_WRITE_GAP"


def test_quote_level_exceptions_become_notes():
    report = build_quote_report(_summary(
        quote_exceptions=[ExceptionDetail(exception_code="CUSTOMER_NOT_FOUND", message="not in IQMS")],
    ))
    assert report.quote_level_notes == ["CUSTOMER_NOT_FOUND: not in IQMS"]


def test_read_gap_codes_match_the_fields_the_webapi_cannot_return():
    assert set(AKA_UNREADABLE_READ_CODES) == set(WEBAPI_UNREADABLE_AKA_FIELDS)


# ---------------------------------------------------------------------------
# Workbook layout
# ---------------------------------------------------------------------------

def test_workbook_has_summary_and_one_sheet_per_quote():
    line = _line(
        aka=AkaResolution(decision="NOT_FOUND"),
        exception_details=[ExceptionDetail(exception_code="AKA_NOT_FOUND")],
    )
    report = build_quote_report(_summary(line))
    data = build_client_workbook([report])

    workbook = load_workbook(BytesIO(data))
    assert workbook.sheetnames == ["Summary", "SAMPLE-Q-001"]

    summary = workbook["Summary"]
    assert summary["A1"].value == "EVCO quote review: summary"
    assert [summary.cell(row=4, column=c).value for c in range(1, 6)] == [
        "Quote no.", "Sheet", "Total rows", "Update required", "No action",
    ]
    assert summary["A5"].value == "SAMPLE-Q-001"
    assert summary["B5"].hyperlink.location == "'SAMPLE-Q-001'!A1"
    assert [summary.cell(row=6, column=c).value for c in (3, 4, 5)] == [1, 1, 0]

    quote = workbook["SAMPLE-Q-001"]
    assert quote["A1"].value == "Quote"
    assert [quote.cell(row=4, column=c).value for c in range(1, 11)] == [
        "Line #", "EVCO part no.", "Customer P/N", "Category", "Rule", "Field",
        "Current value", "Proposed value", "Status", "Comments",
    ]
    assert quote.freeze_panes == "A5"
    assert quote["I5"].value == UPDATE
    assert quote["J5"].value == "AKA_NOT_FOUND; AKA_WRITE_GAP"


def test_sheet_names_are_legal_unique_and_short():
    used = {"Summary"}
    first = _unique_sheet_name("SAMPLE/Q:001", used)
    assert first == "SAMPLE-Q-001"
    used.add(first)
    assert _unique_sheet_name("SAMPLE/Q:001", used) == "SAMPLE-Q-001 (2)"
    assert len(_unique_sheet_name("X" * 80, set())) == 31


def test_client_report_endpoint_returns_xlsx(monkeypatch):
    from app.modules.reporting import router as reporting_router

    monkeypatch.setattr(reporting_router.service, "build_client_report", lambda source, file_name=None: __import__("app.modules.reporting.service", fromlist=["ClientReportFile"]).ClientReportFile(file_name="EVCO_Review_Per_Quote.xlsx", content=b"PK-bytes", sheet_names=[], processing_ids=[]))
    processing_id = str(uuid4())

    response = client.post("/reporting/client-report", json={"processing_ids": [processing_id]})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument")
    assert response.headers["content-disposition"] == 'attachment; filename="EVCO_Review_Per_Quote.xlsx"'
    assert response.content == b"PK-bytes"


def test_client_report_endpoint_requires_at_least_one_quote():
    response = client.post("/reporting/client-report", json={"processing_ids": []})
    assert response.status_code == 400
