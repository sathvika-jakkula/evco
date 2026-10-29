"""
Writes the client review workbook: a Summary sheet plus one sheet per quote,
laid out like the EVCO_Review_Per_Quote workbook.
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import List, Set

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from app.modules.reporting.client_report import NO_ACTION, UPDATE, QuoteReport

NAVY = "1F3864"
UPDATE_FILL = PatternFill("solid", fgColor="FFF2CC")
NO_ACTION_FILL = PatternFill("solid", fgColor="E2EFDA")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
WRAP = Alignment(wrap_text=True, vertical="top")

SUMMARY_HEADERS = ["Quote no.", "Sheet", "Total rows", "Update required", "No action"]
QUOTE_HEADERS = [
    "Line #", "EVCO part no.", "Customer P/N", "Category", "Rule", "Field",
    "Current value", "Proposed value", "Status", "Comments",
]
QUOTE_COLUMN_WIDTHS = [8, 16, 16, 14, 10, 38, 22, 26, 12, 48]
MAX_SHEET_NAME = 31
_ILLEGAL_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def build_client_workbook(reports: List[QuoteReport]) -> bytes:
    return build_client_workbook_with_sheets(reports)[0]


def build_client_workbook_with_sheets(reports: List[QuoteReport]) -> tuple[bytes, List[str]]:
    workbook = Workbook()
    summary = workbook.worksheets[0]
    summary.title = "Summary"

    used_names: Set[str] = {"Summary"}
    sheet_names = []
    for index, report in enumerate(reports, start=1):
        name = _unique_sheet_name(report.quote_number or f"Quote {index}", used_names)
        used_names.add(name)
        sheet_names.append(name)

    _write_summary(summary, reports, sheet_names)
    for report, name in zip(reports, sheet_names):
        _write_quote_sheet(workbook.create_sheet(name), report)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), ["Summary"] + sheet_names


def _unique_sheet_name(raw: str, used: Set[str]) -> str:
    base = _ILLEGAL_SHEET_CHARS.sub("-", raw).strip()[:MAX_SHEET_NAME] or "Quote"
    name, counter = base, 2
    while name in used:
        suffix = f" ({counter})"
        name = base[: MAX_SHEET_NAME - len(suffix)] + suffix
        counter += 1
    return name


def _write_summary(sheet, reports: List[QuoteReport], sheet_names: List[str]) -> None:
    sheet["A1"] = "EVCO quote review: summary"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = "One sheet per quote. Status 'Update required' means a change is needed in IQMS."

    header_row = 4
    for column, title in enumerate(SUMMARY_HEADERS, start=1):
        _style_header(sheet.cell(row=header_row, column=column, value=title))

    row = header_row + 1
    for report, name in zip(reports, sheet_names):
        sheet.cell(row=row, column=1, value=report.quote_number)
        link = sheet.cell(row=row, column=2, value=name)
        link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{name}'!A1")
        link.font = Font(color="0563C1", underline="single")
        sheet.cell(row=row, column=3, value=report.total_rows)
        sheet.cell(row=row, column=4, value=report.update_rows)
        sheet.cell(row=row, column=5, value=report.no_action_rows)
        row += 1

    sheet.cell(row=row, column=1, value="Total").font = Font(bold=True)
    totals = (
        sum(report.total_rows for report in reports),
        sum(report.update_rows for report in reports),
        sum(report.no_action_rows for report in reports),
    )
    # Plain numbers rather than SUM formulas: formulas have no cached value until Excel recalculates.
    for column, total in zip((3, 4, 5), totals):
        sheet.cell(row=row, column=column, value=total).font = Font(bold=True)

    for column, width in zip("ABCDE", [18, 22, 14, 18, 14]):
        sheet.column_dimensions[column].width = width


def _write_quote_sheet(sheet, report: QuoteReport) -> None:
    sheet["A1"] = "Quote"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        f"Customer: {report.customer_number} ({report.customer_name})   "
        f"Processing ID: {report.processing_id}"
    )
    if report.quote_level_notes:
        sheet["A3"] = "Quote-level exceptions: " + "; ".join(report.quote_level_notes)
        sheet["A3"].font = Font(italic=True, color="C00000")

    header_row = 4
    for column, title in enumerate(QUOTE_HEADERS, start=1):
        _style_header(sheet.cell(row=header_row, column=column, value=title))

    for offset, row_data in enumerate(report.rows, start=1):
        row = header_row + offset
        values = [
            row_data.line_number, row_data.evco_part_number, row_data.customer_part_number,
            row_data.category, row_data.rule, row_data.field, row_data.current_value,
            row_data.proposed_value, row_data.status, row_data.comments,
        ]
        fill = UPDATE_FILL if row_data.status == UPDATE else NO_ACTION_FILL if row_data.status == NO_ACTION else None
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.alignment = WRAP
            if fill is not None:
                cell.fill = fill

    for column, width in enumerate(QUOTE_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = f"A{header_row + 1}"
    if report.rows:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(QUOTE_HEADERS))}{header_row + len(report.rows)}"


def _style_header(cell) -> None:
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(vertical="center", wrap_text=True)
